"""In-memory report service with real PDF/JSON/CSV report generation."""
import csv
import io
import json
import os
import uuid
from datetime import datetime, timezone
from typing import Optional

from src.schemas.report import Report, ReportCreate
from src.services import portfolio_service


def _now() -> datetime:
    return datetime.now(timezone.utc)


# In-memory store (keyed by id)
_reports: dict[str, dict] = {}

# Directory for generated report files
REPORTS_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), "tmp_pdf")


_FAKE_DATA_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), "data", "fake")


def _load_fake_portfolio_data() -> list[dict]:
    """Load portfolio holdings from the fake CSV data file as fallback."""
    csv_path = os.path.join(_FAKE_DATA_DIR, "portfolios_holdings.csv")
    if not os.path.exists(csv_path):
        return []

    portfolios: list[dict] = []
    current_portfolio: dict | None = None

    with open(csv_path, "r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            parts = [p.strip() for p in line.split(",")]

            # Detect portfolio header line
            if parts[0] == "portfolio_name":
                continue
            # Detect holdings header line
            if parts[0] == "symbol" and "company_name" in parts:
                continue
            # If line has exactly 2 fields → portfolio definition
            if len(parts) == 2 and not parts[0].isupper():
                current_portfolio = {"name": parts[0], "description": parts[1], "holdings": []}
                portfolios.append(current_portfolio)
                continue
            # Otherwise → holding row (symbol, company_name, quantity, average_price, sector)
            if current_portfolio and len(parts) >= 5:
                try:
                    qty = float(parts[2])
                    price = float(parts[3])
                    current_portfolio["holdings"].append({
                        "symbol": parts[0],
                        "company": parts[1],
                        "sector": parts[4] if len(parts) > 4 else "N/A",
                        "quantity": qty,
                        "avg_price": price,
                        "market_value": round(qty * price, 2),
                    })
                except (ValueError, IndexError):
                    continue

    return portfolios


def _load_fake_commodity_data() -> list[dict]:
    """Load commodity prices from the fake CSV data file."""
    csv_path = os.path.join(_FAKE_DATA_DIR, "commodity_prices.csv")
    if not os.path.exists(csv_path):
        return []

    commodities: list[dict] = []
    with open(csv_path, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row in reader:
            try:
                commodities.append({
                    "commodity": row["commodity_type"],
                    "date": row["date"],
                    "price": float(row["price"]),
                    "unit": row["unit"],
                })
            except (KeyError, ValueError):
                continue
    return commodities


def _get_portfolio_data(portfolio_id: Optional[str]) -> dict:
    """Gather portfolio holdings data for the report.
    Falls back to fake CSV data when in-memory store is empty."""
    holdings_data = []
    portfolio_name = "All Portfolios"
    total_value = 0.0

    # Try in-memory store first
    if portfolio_id:
        portfolio = portfolio_service.get_portfolio(portfolio_id)
        if portfolio:
            portfolio_name = portfolio.name
        result = portfolio_service.list_holdings(portfolio_id)
        if result:
            holdings_list, _ = result
            for h in holdings_list:
                value = h.quantity * h.average_price
                total_value += value
                holdings_data.append({
                    "symbol": h.symbol,
                    "company": h.company_name,
                    "sector": h.sector or "N/A",
                    "quantity": h.quantity,
                    "avg_price": h.average_price,
                    "market_value": round(value, 2),
                })
    else:
        portfolios, _ = portfolio_service.list_portfolios(skip=0, limit=100)
        for p in portfolios:
            result = portfolio_service.list_holdings(p.id)
            if result:
                holdings_list, _ = result
                for h in holdings_list:
                    value = h.quantity * h.average_price
                    total_value += value
                    holdings_data.append({
                        "symbol": h.symbol,
                        "company": h.company_name,
                        "sector": h.sector or "N/A",
                        "quantity": h.quantity,
                        "avg_price": h.average_price,
                        "market_value": round(value, 2),
                    })

    # Fallback: load from fake CSV if no in-memory data found
    if not holdings_data:
        fake_portfolios = _load_fake_portfolio_data()
        if fake_portfolios:
            # If portfolio_id matches a portfolio name, use that; otherwise use first
            target = None
            for fp in fake_portfolios:
                if portfolio_id and portfolio_id.lower() in fp["name"].lower():
                    target = fp
                    break
            if not target:
                target = fake_portfolios[0]  # Default to first portfolio
                if not portfolio_id:
                    # Use all portfolios combined
                    portfolio_name = "All Portfolios (Sample Data)"
                    for fp in fake_portfolios:
                        for h in fp["holdings"]:
                            # Avoid duplicate symbols across portfolios
                            if not any(eh["symbol"] == h["symbol"] for eh in holdings_data):
                                holdings_data.append(h)
                                total_value += h["market_value"]
                    target = None  # Already processed

            if target:
                portfolio_name = f"{target['name']} (Sample Data)"
                holdings_data = target["holdings"]
                total_value = sum(h["market_value"] for h in holdings_data)

    # Compute sector allocation
    sector_map: dict[str, float] = {}
    for h in holdings_data:
        sector_map[h["sector"]] = sector_map.get(h["sector"], 0) + h["market_value"]

    sector_allocation = []
    for sector, value in sorted(sector_map.items(), key=lambda x: -x[1]):
        pct = (value / total_value * 100) if total_value > 0 else 0
        sector_allocation.append({"sector": sector, "value": round(value, 2), "pct": round(pct, 1)})

    # Load commodity data for market context
    commodities = _load_fake_commodity_data()
    # Get latest prices per commodity
    latest_commodities: dict[str, dict] = {}
    for c in commodities:
        existing = latest_commodities.get(c["commodity"])
        if not existing or c["date"] > existing["date"]:
            latest_commodities[c["commodity"]] = c

    return {
        "portfolio_name": portfolio_name,
        "holdings": holdings_data,
        "total_value": round(total_value, 2),
        "sector_allocation": sector_allocation,
        "commodities": list(latest_commodities.values()),
    }


def _generate_pdf_report(report_name: str, report_type: str, data: dict, file_path: str) -> None:
    """Generate a real PDF report using reportlab."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    doc = SimpleDocTemplate(file_path, pagesize=A4)
    styles = getSampleStyleSheet()
    story = []

    # Title
    title_style = ParagraphStyle("ReportTitle", parent=styles["Title"], fontSize=20, spaceAfter=20, textColor=colors.HexColor("#1a237e"))
    story.append(Paragraph(report_name, title_style))
    story.append(Spacer(1, 0.1 * inch))

    # Subtitle
    subtitle_style = ParagraphStyle("Subtitle", parent=styles["Normal"], fontSize=12, textColor=colors.grey, spaceAfter=20)
    story.append(Paragraph(f"Report Type: {report_type.replace('_', ' ').title()}", subtitle_style))
    story.append(Paragraph(f"Portfolio: {data['portfolio_name']}", subtitle_style))
    story.append(Paragraph(f"Generated: {_now().strftime('%Y-%m-%d %H:%M UTC')}", subtitle_style))
    story.append(Spacer(1, 0.3 * inch))

    # Portfolio Summary
    section_style = ParagraphStyle("SectionTitle", parent=styles["Heading2"], fontSize=14, textColor=colors.HexColor("#283593"), spaceAfter=10)
    story.append(Paragraph("Portfolio Summary", section_style))
    story.append(Paragraph(f"Total Portfolio Value: ${data['total_value']:,.2f}", styles["Normal"]))
    story.append(Paragraph(f"Number of Holdings: {len(data['holdings'])}", styles["Normal"]))
    story.append(Spacer(1, 0.2 * inch))

    # Holdings Table
    if data["holdings"]:
        story.append(Paragraph("Holdings Breakdown", section_style))
        table_data = [["Symbol", "Company", "Sector", "Qty", "Avg Price", "Market Value"]]
        for h in data["holdings"]:
            table_data.append([
                h["symbol"],
                h["company"][:25],
                h["sector"].replace("_", " ").title(),
                str(h["quantity"]),
                f"${h['avg_price']:,.2f}",
                f"${h['market_value']:,.2f}",
            ])

        t = Table(table_data, colWidths=[60, 130, 100, 50, 70, 90])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1a237e")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 9),
            ("FONTSIZE", (0, 1), (-1, -1), 8),
            ("ALIGN", (3, 0), (-1, -1), "RIGHT"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f5f5")]),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(t)
        story.append(Spacer(1, 0.3 * inch))

    # Sector Allocation Table
    if data["sector_allocation"]:
        story.append(Paragraph("Sector Allocation", section_style))
        sec_data = [["Sector", "Market Value", "Weight (%)"]]
        for s in data["sector_allocation"]:
            sec_data.append([
                s["sector"].replace("_", " ").title(),
                f"${s['value']:,.2f}",
                f"{s['pct']}%",
            ])

        st = Table(sec_data, colWidths=[150, 120, 80])
        st.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#283593")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#e8eaf6")]),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(st)
        story.append(Spacer(1, 0.3 * inch))

    # Commodities / Market Factors
    if "commodities" in data and data["commodities"]:
        story.append(Paragraph("Market Factors (Commodities)", section_style))
        comm_data = [["Commodity", "Latest Price", "Date"]]
        for c in data["commodities"]:
            comm_data.append([
                c["commodity"].replace("_", " ").title(),
                f"${c['price']:,.2f} {c['unit']}",
                c["date"],
            ])

        ct = Table(comm_data, colWidths=[150, 120, 80])
        ct.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#283593")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#e8eaf6")]),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(ct)

    # If no holdings data
    if not data["holdings"]:
        story.append(Paragraph("No holdings found for this portfolio. Create a portfolio with holdings first to see report data.", styles["Normal"]))

    # Footer
    story.append(Spacer(1, 0.5 * inch))
    footer_style = ParagraphStyle("Footer", parent=styles["Normal"], fontSize=8, textColor=colors.grey)
    story.append(Paragraph("Generated by PortfolioQ – Autonomous Market & Portfolio Scenario Analyst", footer_style))

    doc.build(story)


def _generate_json_report(report_name: str, report_type: str, data: dict, file_path: str) -> None:
    """Generate a JSON report."""
    report = {
        "report_name": report_name,
        "report_type": report_type,
        "generated_at": _now().isoformat(),
        "portfolio": data["portfolio_name"],
        "total_value": data["total_value"],
        "holdings_count": len(data["holdings"]),
        "holdings": data["holdings"],
        "sector_allocation": data["sector_allocation"],
        "market_factors": data.get("commodities", []),
    }
    with open(file_path, "w", encoding="utf-8") as f:
        json.dump(report, f, indent=2)


def _generate_csv_report(report_name: str, report_type: str, data: dict, file_path: str) -> None:
    """Generate a CSV report."""
    with open(file_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["Report", report_name])
        writer.writerow(["Type", report_type])
        writer.writerow(["Generated", _now().isoformat()])
        writer.writerow(["Portfolio", data["portfolio_name"]])
        writer.writerow(["Total Value", data["total_value"]])
        writer.writerow([])
        writer.writerow(["Symbol", "Company", "Sector", "Quantity", "Avg Price", "Market Value"])
        for h in data["holdings"]:
            writer.writerow([h["symbol"], h["company"], h["sector"], h["quantity"], h["avg_price"], h["market_value"]])
        writer.writerow([])
        writer.writerow(["Sector", "Value", "Weight (%)"])
        for s in data["sector_allocation"]:
            writer.writerow([s["sector"], s["value"], s["pct"]])
        if data.get("commodities"):
            writer.writerow([])
            writer.writerow(["Market Factors"])
            writer.writerow(["Commodity", "Price", "Unit", "Date"])
            for c in data["commodities"]:
                writer.writerow([c["commodity"], c["price"], c["unit"], c["date"]])


def _generate_excel_report(report_name: str, report_type: str, data: dict, file_path: str) -> None:
    """Generate an Excel report using openpyxl."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Portfolio Report"

    header_fill = PatternFill(start_color="1a237e", end_color="1a237e", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=10)
    title_font = Font(bold=True, size=14, color="1a237e")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    # Title
    ws["A1"] = report_name
    ws["A1"].font = title_font
    ws["A2"] = f"Type: {report_type.replace('_', ' ').title()}"
    ws["A3"] = f"Portfolio: {data['portfolio_name']}"
    ws["A4"] = f"Generated: {_now().strftime('%Y-%m-%d %H:%M UTC')}"
    ws["A5"] = f"Total Value: ${data['total_value']:,.2f}"

    # Holdings header
    row = 7
    headers = ["Symbol", "Company", "Sector", "Quantity", "Avg Price", "Market Value"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border

    for h in data["holdings"]:
        row += 1
        vals = [h["symbol"], h["company"], h["sector"].replace("_", " ").title(),
                h["quantity"], h["avg_price"], h["market_value"]]
        for col, v in enumerate(vals, 1):
            cell = ws.cell(row=row, column=col, value=v)
            cell.border = thin_border

    # Sector allocation
    row += 2
    ws.cell(row=row, column=1, value="Sector Allocation").font = Font(bold=True, size=12, color="283593")
    row += 1
    for col, h in enumerate(["Sector", "Value", "Weight (%)"], 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.fill = PatternFill(start_color="283593", end_color="283593", fill_type="solid")
        cell.font = header_font
        cell.border = thin_border

    for s in data["sector_allocation"]:
        row += 1
        ws.cell(row=row, column=1, value=s["sector"].replace("_", " ").title()).border = thin_border
        ws.cell(row=row, column=2, value=s["value"]).border = thin_border
        ws.cell(row=row, column=3, value=s["pct"]).border = thin_border

    # Commodities
    if data.get("commodities"):
        row += 2
        ws.cell(row=row, column=1, value="Market Factors (Commodities)").font = Font(bold=True, size=12, color="283593")
        row += 1
        for col, h in enumerate(["Commodity", "Price", "Unit", "Date"], 1):
            cell = ws.cell(row=row, column=col, value=h)
            cell.fill = PatternFill(start_color="283593", end_color="283593", fill_type="solid")
            cell.font = header_font
            cell.border = thin_border

        for c in data["commodities"]:
            row += 1
            ws.cell(row=row, column=1, value=c["commodity"].replace("_", " ").title()).border = thin_border
            ws.cell(row=row, column=2, value=c["price"]).border = thin_border
            ws.cell(row=row, column=3, value=c["unit"]).border = thin_border
            ws.cell(row=row, column=4, value=c["date"]).border = thin_border

    # Column widths
    ws.column_dimensions["A"].width = 15
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 15

    wb.save(file_path)


def list_reports(
    skip: int = 0,
    limit: int = 20,
) -> tuple[list[Report], int]:
    """Return (reports slice, total count)."""
    items = list(_reports.values())
    total = len(items)
    items.sort(key=lambda r: r["generated_at"], reverse=True)
    slice_items = items[skip : skip + limit]
    reports = [Report(**r) for r in slice_items]
    return reports, total


def get_report(report_id: str) -> Optional[Report]:
    """Get a report by ID or None if not found."""
    r = _reports.get(report_id)
    if not r:
        return None
    return Report(**r)


def create_report(payload: ReportCreate) -> Report:
    """Create a pending report record (for testing only)."""
    rid = str(uuid.uuid4())
    now = _now()
    r = {
        "id": rid,
        "name": payload.name,
        "type": payload.type,
        "format": payload.format,
        "portfolio_id": payload.portfolio_id,
        "scenario_id": payload.scenario_id,
        "file_path": None,
        "download_url": None,
        "status": "pending",
        "generated_at": now,
    }
    _reports[rid] = r
    return Report(**r)


def generate_report(payload: ReportCreate) -> Report:
    """Generate a report synchronously. Returns a completed Report with file_path set."""
    rid = str(uuid.uuid4())
    now = _now()
    fmt = payload.format.value if hasattr(payload.format, "value") else payload.format

    os.makedirs(REPORTS_DIR, exist_ok=True)

    # Determine file extension
    ext_map = {"pdf": "pdf", "excel": "xlsx", "json": "json", "csv": "csv"}
    ext = ext_map.get(fmt, "txt")
    filename = f"report_{rid}.{ext}"
    file_path = os.path.join(REPORTS_DIR, filename)

    # Gather real portfolio data
    data = _get_portfolio_data(payload.portfolio_id)

    # Generate the actual file
    try:
        if fmt == "pdf":
            _generate_pdf_report(payload.name, payload.type, data, file_path)
        elif fmt == "json":
            _generate_json_report(payload.name, payload.type, data, file_path)
        elif fmt == "csv":
            _generate_csv_report(payload.name, payload.type, data, file_path)
        elif fmt == "excel":
            _generate_excel_report(payload.name, payload.type, data, file_path)
        else:
            # Fallback: plain text
            with open(file_path, "w") as f:
                f.write(f"Report: {payload.name}\nType: {payload.type}\n")
        status = "completed"
    except Exception as e:
        status = "failed"
        file_path = None

    r = {
        "id": rid,
        "name": payload.name,
        "type": payload.type,
        "format": payload.format,
        "portfolio_id": payload.portfolio_id,
        "scenario_id": payload.scenario_id,
        "file_path": file_path,
        "download_url": f"/api/v1/reports/{rid}/download" if status == "completed" else None,
        "status": status,
        "generated_at": now,
    }
    _reports[rid] = r
    return Report(**r)


def clear_all() -> None:
    """Clear all in-memory data. For testing only."""
    _reports.clear()
