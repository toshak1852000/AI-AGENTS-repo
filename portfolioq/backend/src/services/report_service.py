"""Report generation service: PDF, Excel, and JSON board-ready reports."""
import io
import json
import logging
import os
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from sqlalchemy.orm import Session

from src.models.report import Report
from src.ml.llm_service import generate_scenario_narrative

logger = logging.getLogger(__name__)
REPORTS_DIR = Path(os.getenv("REPORTS_DIR", "/app/reports"))


def generate_report(
    db: Session,
    scenario: dict[str, Any],
    portfolio: dict[str, Any],
    exposure_result: dict[str, Any],
    risk_result: dict[str, Any],
    run_id: str,
    report_format: str = "json",
) -> dict[str, Any]:
    """Generate and persist a board-ready scenario analysis report."""
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    report_id = str(uuid.uuid4())
    timestamp = datetime.now(timezone.utc)
    report_name = (
        f"Scenario_Report_{scenario.get('name','Unknown')}_"
        f"{timestamp.strftime('%Y%m%d_%H%M%S')}"
    )

    narrative = generate_scenario_narrative(scenario, portfolio, exposure_result, risk_result)
    recommendations = _extract_recommendations(narrative)

    summary = {
        "report_id": report_id,
        "run_id": run_id,
        "generated_at": timestamp.isoformat(),
        "scenario": {"id": scenario.get("id"), "name": scenario.get("name"), "type": scenario.get("type")},
        "portfolio": {"id": portfolio.get("id"), "name": portfolio.get("name")},
        "executive_summary": narrative,
        "risk_assessment": {
            "risk_level": risk_result.get("risk_level"),
            "risk_score": risk_result.get("risk_score"),
            "pl_impact": risk_result.get("pl_impact"),
            "pl_impact_pct": risk_result.get("pl_impact_pct"),
        },
        "exposure_summary": {
            "total_market_value": exposure_result.get("total_exposure"),
            "sensitivity_score": exposure_result.get("sensitivity_score"),
            "top_sectors": (exposure_result.get("sector_exposures") or [])[:5],
        },
        "recommendations": recommendations,
        "top_holdings_at_risk": (risk_result.get("prioritized_holdings") or [])[:10],
    }

    file_path: str | None = None
    try:
        if report_format == "json":
            file_path = _write_json(report_id, summary)
        elif report_format == "pdf":
            file_path = _write_pdf(report_id, summary)
        elif report_format == "excel":
            file_path = _write_excel(report_id, summary, exposure_result, risk_result)
    except Exception as exc:
        logger.warning("Could not write report file (%s): %s", report_format, exc)

    db_report = Report(
        id=report_id,
        name=report_name,
        type="scenario_analysis",
        format=report_format,
        portfolio_id=portfolio.get("id") or None,
        scenario_id=scenario.get("id") or None,
        run_id=run_id,
        file_path=file_path,
        summary=summary,
    )
    try:
        db.add(db_report)
        db.commit()
    except Exception as exc:
        logger.warning("Could not persist report: %s", exc)
        db.rollback()

    return summary


def _extract_recommendations(narrative: str) -> list[str]:
    recs = []
    lines = narrative.splitlines()
    in_recs = False
    for line in lines:
        if "RECOMMENDATION" in line.upper():
            in_recs = True
            continue
        if in_recs and line.strip().startswith(("1.", "2.", "3.", "4.", "5.", "-", "•", "*")):
            recs.append(line.strip().lstrip("1234567890.-•* "))
        elif in_recs and line.strip() and not any(c.isupper() for c in line[:4]):
            pass
        elif in_recs and line.strip() == "":
            pass
        elif in_recs and line.strip().isupper():
            break
    return recs[:5] or ["No specific recommendations at this time."]


def _write_json(report_id: str, summary: dict) -> str:
    path = REPORTS_DIR / f"{report_id}.json"
    path.write_text(json.dumps(summary, indent=2, default=str))
    return str(path)


def _write_pdf(report_id: str, summary: dict) -> str:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

    path = str(REPORTS_DIR / f"{report_id}.pdf")
    doc = SimpleDocTemplate(path, pagesize=A4, topMargin=2*cm, bottomMargin=2*cm)
    styles = getSampleStyleSheet()
    heading_style = ParagraphStyle("heading", parent=styles["Heading1"], fontSize=16, spaceAfter=12)
    subheading = ParagraphStyle("subhead", parent=styles["Heading2"], fontSize=12, spaceAfter=8)
    body = styles["BodyText"]

    story = []
    story.append(Paragraph("PortfolioQ — Scenario Analysis Report", heading_style))
    story.append(Paragraph(f"Report ID: {report_id}", body))
    story.append(Paragraph(f"Generated: {summary.get('generated_at','')}", body))
    story.append(Spacer(1, 0.5*cm))

    story.append(Paragraph("Scenario Details", subheading))
    sc = summary.get("scenario", {})
    story.append(Paragraph(f"Name: {sc.get('name','N/A')} | Type: {sc.get('type','N/A')}", body))
    story.append(Spacer(1, 0.3*cm))

    story.append(Paragraph("Risk Assessment", subheading))
    ra = summary.get("risk_assessment", {})
    story.append(Paragraph(
        f"Risk Level: <b>{str(ra.get('risk_level','')).upper()}</b>  |  "
        f"Risk Score: {ra.get('risk_score', 0):.2f}  |  "
        f"P&L Impact: ${ra.get('pl_impact', 0):,.0f}  ({ra.get('pl_impact_pct', 0):.2f}%)", body))
    story.append(Spacer(1, 0.3*cm))

    story.append(Paragraph("Executive Summary", subheading))
    for line in summary.get("executive_summary", "").splitlines():
        if line.strip():
            story.append(Paragraph(line.strip(), body))
    story.append(Spacer(1, 0.3*cm))

    # Top sectors table
    sectors = summary.get("exposure_summary", {}).get("top_sectors", [])
    if sectors:
        story.append(Paragraph("Sector Exposure", subheading))
        tbl_data = [["Sector", "Market Value ($)", "Exposure %", "P&L Impact ($)"]]
        for s in sectors[:8]:
            tbl_data.append([
                s.get("sector", ""), f"${s.get('market_value', 0):,.0f}",
                f"{s.get('exposure_pct', 0):.1%}", f"${s.get('scenario_pl_impact', 0):,.0f}",
            ])
        tbl = Table(tbl_data, colWidths=[5*cm, 4*cm, 3*cm, 4*cm])
        tbl.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e40af")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f5f9")]),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
        ]))
        story.append(tbl)
        story.append(Spacer(1, 0.3*cm))

    # Recommendations
    recs = summary.get("recommendations", [])
    if recs:
        story.append(Paragraph("Strategic Recommendations", subheading))
        for i, rec in enumerate(recs, 1):
            story.append(Paragraph(f"{i}. {rec}", body))

    doc.build(story)
    return path


def _write_excel(report_id: str, summary: dict, exposure_result: dict, risk_result: dict) -> str:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    path = str(REPORTS_DIR / f"{report_id}.xlsx")
    wb = openpyxl.Workbook()

    # Summary sheet
    ws = wb.active
    ws.title = "Summary"
    header_fill = PatternFill("solid", fgColor="1e40af")
    header_font = Font(bold=True, color="FFFFFF", size=11)

    ws.append(["PortfolioQ — Scenario Analysis Report"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(["Report ID", report_id])
    ws.append(["Generated At", summary.get("generated_at", "")])
    ws.append(["Scenario", summary.get("scenario", {}).get("name", "")])
    ws.append(["Portfolio", summary.get("portfolio", {}).get("name", "")])
    ws.append([])
    ws.append(["Risk Level", summary.get("risk_assessment", {}).get("risk_level", "").upper()])
    ws.append(["Risk Score", summary.get("risk_assessment", {}).get("risk_score", 0)])
    ws.append(["P&L Impact ($)", summary.get("risk_assessment", {}).get("pl_impact", 0)])
    ws.append(["P&L Impact (%)", summary.get("risk_assessment", {}).get("pl_impact_pct", 0)])

    # Company Exposure sheet
    ws2 = wb.create_sheet("Company Exposure")
    headers = ["Symbol", "Market Value ($)", "Exposure %", "P&L Impact ($)", "Risk Score"]
    ws2.append(headers)
    for cell in ws2[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    for ce in exposure_result.get("company_exposures", []):
        ws2.append([
            ce.get("symbol"), round(ce.get("market_value", 0), 2),
            ce.get("exposure_pct", 0), round(ce.get("scenario_pl_impact", 0), 2),
            round(ce.get("risk_score", 0), 4),
        ])
    ws2.column_dimensions["A"].width = 12

    # Sector Exposure sheet
    ws3 = wb.create_sheet("Sector Exposure")
    ws3.append(["Sector", "Market Value ($)", "Exposure %", "P&L Impact ($)", "Holdings"])
    for cell in ws3[1]:
        cell.fill = header_fill
        cell.font = header_font
    for se in exposure_result.get("sector_exposures", []):
        ws3.append([
            se.get("sector"), round(se.get("market_value", 0), 2),
            se.get("exposure_pct", 0), round(se.get("scenario_pl_impact", 0), 2),
            ", ".join(se.get("symbols", [])),
        ])

    # Prioritized Holdings sheet
    ws4 = wb.create_sheet("Prioritized Holdings")
    ws4.append(["Symbol", "Risk Level", "Risk Score", "P&L Impact ($)", "Signal", "Exposure %"])
    for cell in ws4[1]:
        cell.fill = header_fill
        cell.font = header_font
    for h in (risk_result.get("prioritized_holdings") or [])[:30]:
        ws4.append([
            h.get("symbol"), h.get("risk_level", ""), round(h.get("risk_score", 0), 4),
            round(h.get("scenario_pl_impact", 0), 2), h.get("signal", ""),
            h.get("exposure_pct", 0),
        ])

    wb.save(path)
    return path


def get_report_file(db: Session, report_id: str) -> tuple[str | None, str | None]:
    """Return (file_path, format) for a report."""
    row = db.query(Report).filter(Report.id == report_id).first()
    if not row:
        return None, None
    return row.file_path, row.format
